"""Checagens novas da auditoria externa: siglas, MeSH, IA nos Métodos, números, PRISMA 2020 e referências."""

import json

import pytest

from artigos_v2.auditoria import runner
from artigos_v2.auditoria import skills as S
from artigos_v2.auditoria.leitor import ler
from artigos_v2.auditoria.runner import auditar
from artigos_v2.auditoria.verificacao import Verificador

PRISMA = ("Page MJ, McKenzie JE, Bossuyt PM, Boutron I, Hoffmann TC, Mulrow CD, et al. The PRISMA 2020 statement: "
          "an updated guideline for reporting systematic reviews. BMJ. 2021;372:n71. doi:10.1136/bmj.n71")

REFS_BOAS = [
    PRISMA,
    "Malta DC, Gonçalves RPF, Machado IE, Freitas MIF, Azeredo C, Szwarcwald CL. Prevalence of arterial "
    "hypertension according to different diagnostic criteria, National Health Survey. Rev Bras Epidemiol. "
    "2018;21(Suppl 1):e180021. doi:10.1590/1980-549720180021.supl.1",
    "Brasil. Ministério da Saúde. Estratégias para o cuidado da pessoa com doença crônica: hipertensão arterial "
    "sistêmica. Brasília: Ministério da Saúde; 2013.",
    "World Health Organization. Hypertension [Internet]. Geneva: WHO; 2023 [cited 2024 Jan 10]. Available from: "
    "https://www.who.int/news-room/fact-sheets/detail/hypertension",
    "de Souza AB, Lima ÂM, et al. Adesão ao tratamento anti-hipertensivo na atenção primária. Cad Saude Publica. "
    "2020;36(5):e00123419. PMID: 32401234",
    "Whelton PK, Carey RM, Aronow WS. 2017 guideline for high blood pressure in adults. Hypertension. "
    "2018 Jun;71(6):1269-324. doi:10.1161/HYP.0000000000000065",
    "SILVA, A. B.; SOUZA, C. D. Hipertensão na atenção primária. Revista X, v. 1, n. 2, p. 10-20, 2020. "
    "DOI: https://doi.org/10.1590/abc.123.",
]

REFS_RUINS = [
    "Altamirano, Pérez JL, Gómez R. Predictive factors of difficult laparoscopic cholecystectomy. Surg Endosc. "
    "2019;33(2).",
    "Pérez JL, Gómez R. Colecistectomia difícil em hospital de ensino. Rev. Esp. Enferm. Dig. 2019;111(5):300-5..",
    "Lima EF, Costa GH. Hypertension care in Brazil. Arq Bras Cardiol. 2019;112(3):200-8. doi:10.5935/abc.2019. "
    "Disponível em: https://www.scielo.br/j/abc/a/xyz",
    "Souza MA,. Guia de hipertensão para a APS. Disponível em: https://exemplo.org/guia",
]


def _item(resultado, trecho):
    return next(i for i in resultado["itens"] if trecho in i["criterio"])


def _artigo(tmp_path, nome="a.md", intro="", metodos="", resultados="", discussao="", conclusao="",
            refs=(), legendas="", tabela_metodos="", tabela_resultados=""):
    partes = ["# Título", "## Introdução", intro or "Texto.", "## Métodos", metodos or "Texto."]
    if tabela_metodos:
        partes += ["Quadro 1 – Estratégias de busca", tabela_metodos, "Fonte: elaborado pelos autores (2026)."]
    partes += ["## Resultados", resultados or "Texto."]
    if tabela_resultados:
        partes.append(tabela_resultados)
    if legendas:
        partes.append(legendas)
    partes += ["## Discussão", discussao or "Texto.", "## Conclusão", conclusao or "Texto.", "## Referências",
               "\n".join(f"{i}. {r}" for i, r in enumerate(refs, 1)) or "1. Silva AB. Estudo. Rev X. 2020;1:1."]
    caminho = tmp_path / nome
    caminho.write_text("\n\n".join(partes) + "\n", encoding="utf-8")
    return caminho


# ------------------------------------------------------------------ leitor

def test_leitor_separa_subtitulos_e_conta_citacao_na_fonte(tmp_path):
    caminho = _artigo(tmp_path, metodos="A busca seguiu o PRISMA 2020 [1].\n\n### 2.1 Busca na BVS\n\nOutro parágrafo.",
                      legendas="Figura 1 – Fluxograma\n\nFonte: adaptado do modelo PRISMA 2020 [2].",
                      refs=["Silva AB. Um. Rev X. 2020;1:1.", PRISMA])
    artigo = ler(caminho)
    assert artigo.subtitulos == ["2.1 Busca na BVS"]
    assert not any("BVS" in p for p in artigo.secoes["metodos"])
    assert [2] in artigo.citacoes


def test_leitor_docx_citacao_sobrescrita_na_fonte_e_subtitulo(tmp_path):
    from docx import Document

    documento = Document()
    documento.add_heading("1 INTRODUÇÃO", level=1)
    documento.add_paragraph("A Atenção Primária à Saúde (APS) organiza o cuidado.")
    documento.add_heading("2 MÉTODOS", level=1)
    documento.add_heading("2.1 Busca na BVS", level=2)
    documento.add_paragraph("Os autores buscaram no PubMed.")
    fonte = documento.add_paragraph("Fonte: adaptado do modelo PRISMA 2020")
    fonte.add_run("1").font.superscript = True
    fonte.add_run(".")
    documento.add_heading("REFERÊNCIAS", level=1)
    documento.add_paragraph("1. " + PRISMA)
    caminho = tmp_path / "a.docx"
    documento.save(caminho)

    artigo = ler(caminho)
    assert artigo.subtitulos == ["2.1 Busca na BVS"] and artigo.secoes["metodos"] == ["Os autores buscaram no PubMed."]
    assert artigo.legendas == ["Fonte: adaptado do modelo PRISMA 2020[^1]."] and artigo.citacoes == [[1]]
    assert _item(S.skill_resultados(artigo), "PRISMA 2020 citado")["status"] == "ok"
    assert _item(S.skill_estrutura_introducao(artigo), "Siglas")["status"] == "ok"


# ------------------------------------------------------------------ siglas

def test_siglas_definidas_na_primeira_ocorrencia(tmp_path):
    boa = _artigo(tmp_path, "boa.md",
                  intro="A Atenção Primária à Saúde (APS) organiza o cuidado no Sistema Único de Saúde (SUS). "
                        "No século XX, a APS cresceu; o DNA não precisa de definição.",
                  metodos="Usaram-se Descritores em Ciências da Saúde (DeCS) e Medical Subject Headings (MeSH), "
                          "combinados por AND e OR. A estratégia foi (Hypertension[MeSH Terms] AND LANG:\"por\").\n\n"
                          "### Busca na BVS\n\nO SUS e as Unidades Básicas de Saúde (UBSs) foram o cenário.",
                  discussao="As UBS e a APS seguem centrais no SUS.")
    ruim = _artigo(tmp_path, "ruim.md",
                   intro="A APS organiza o cuidado. A Atenção Primária à Saúde (APS) é central.",
                   metodos="Usaram-se descritores DeCS e MeSH.",
                   discussao="A COVID-19 afetou o SUS.")
    ok = S.skill_estrutura_introducao(ler(boa))
    assert _item(ok, "Siglas definidas")["status"] == "ok", _item(ok, "Siglas definidas")
    item = _item(S.skill_estrutura_introducao(ler(ruim)), "Siglas definidas")
    assert item["status"] == "falha"
    for sigla in ("APS", "DeCS", "MeSH", "COVID-19", "SUS"):
        assert sigla in item["evidencia"]


# ------------------------------------------------------------------ Métodos

class _VerificadorMesh(Verificador):
    """MeSH simulado: sem rede."""

    VALIDOS = {"hypertension": "Hypertension", "primary health care": "Primary Health Care",
               "cholecystectomy, laparoscopic": "Cholecystectomy, Laparoscopic"}
    SINONIMOS = {"high blood pressure": "Hypertension"}

    def __init__(self):
        super().__init__()
        self.consultados = []

    def verificar_descritor_mesh(self, termo):
        self.consultados.append(termo)
        chave = termo.lower()
        if chave in self.VALIDOS:
            return {"existe": True, "descritor": self.VALIDOS[chave]}
        return {"existe": False, "descritor": self.SINONIMOS.get(chave, "")}

    def contar_pubmed(self, estrategia):
        return 40


def test_extrai_descritores_mesh_de_varias_marcacoes():
    estrategia = ('("Predictive Factors"[MeSH Terms] OR Hypertension[mh] OR Cholecystectomy, Laparoscopic[mh]) '
                  'AND "Hypertension/therapy"[majr] AND "high blood pressure"[MeSH] AND risk[tiab] '
                  'AND primary health care[mh:noexp]')
    assert S.descritores_mesh(estrategia) == ["Predictive Factors", "Hypertension", "Cholecystectomy, Laparoscopic",
                                              "high blood pressure", "primary health care"]
    # estratégia escrita na prosa dos Métodos, sem aspas nem parênteses
    prosa = ("A estratégia no PubMed foi Hypertension[mh] AND Cholecystectomy, Laparoscopic[majr]. Depois, "
             "Diabetes Mellitus, Type 2[MeSH Terms].")
    assert S.descritores_mesh(prosa) == ["Hypertension", "Cholecystectomy, Laparoscopic", "Diabetes Mellitus, Type 2"]


def test_descritores_mesh_verificados_no_quadro(tmp_path):
    quadro = ("| Base | Estratégia | Data | Registros (n) |\n|---|---|---|---|\n"
              "| PubMed | (\"Predictive Factors\"[MeSH Terms] OR \"High Blood Pressure\"[mh]) AND "
              "Hypertension[MeSH Terms] | 26/09/2026 | 40 |")
    caminho = _artigo(tmp_path, metodos="A busca foi feita no PubMed.", tabela_metodos=quadro)
    verificador = _VerificadorMesh()
    item = _item(S.skill_metodos(ler(caminho), verificador=verificador), "Descritores MeSH")
    assert item["status"] == "falha"
    assert "Predictive Factors" in item["evidencia"] and '"Hypertension"' in item["evidencia"]
    assert verificador.consultados == ["Predictive Factors", "High Blood Pressure", "Hypertension"]

    bom = quadro.replace('"Predictive Factors"[MeSH Terms] OR "High Blood Pressure"[mh]', '"Primary Health Care"[mh]')
    caminho = _artigo(tmp_path, "bom.md", metodos="A busca foi feita no PubMed.", tabela_metodos=bom)
    assert _item(S.skill_metodos(ler(caminho), verificador=_VerificadorMesh()), "Descritores MeSH")["status"] == "ok"
    # offline e rede fora do ar: não verificável, nunca falha
    assert _item(S.skill_metodos(ler(caminho)), "Descritores MeSH")["status"] == "nao_verificavel"

    class _SemRede(_VerificadorMesh):
        def verificar_descritor_mesh(self, termo):
            raise OSError("sem rede")

    assert _item(S.skill_metodos(ler(caminho), verificador=_SemRede()), "Descritores MeSH")["status"] == "nao_verificavel"


def test_descritores_mesh_no_texto_dos_metodos(tmp_path):
    caminho = _artigo(tmp_path, metodos='A estratégia no PubMed foi ("Predictive Factors"[MeSH] AND '
                                        'cholecystectomy[tiab]).')
    item = _item(S.skill_metodos(ler(caminho), verificador=_VerificadorMesh()), "Descritores MeSH")
    assert item["status"] == "falha" and "Predictive Factors" in item["evidencia"]


def test_auditar_online_usa_o_verificador_de_mesh(tmp_path, monkeypatch):
    monkeypatch.setattr(runner, "Verificador", lambda *_a, **_k: _VerificadorMesh())
    monkeypatch.setattr(_VerificadorMesh, "verificar_referencia",
                        lambda self, ref: {"status": "verificada", "fonte": "teste", "titulo_base": ""})
    caminho = _artigo(tmp_path, metodos='No PubMed: ("Predictive Factors"[MeSH Terms] AND Hypertension[mh]).')
    rel = auditar(caminho)
    metodos = next(s for s in rel["skills"] if s["id"] == "metodos")
    assert _item(metodos, "Descritores MeSH")["status"] == "falha"


@pytest.mark.parametrize("texto, status", [
    ("Dois autores fizeram a triagem de forma independente; divergências foram resolvidas pelo orientador.", "ok"),
    ("Os revisores extraíram os dados.", "ok"),
    ("Foram buscados artigos no PubMed.", "parcial"),
])
def test_metodos_descrevem_processo_dos_autores(tmp_path, texto, status):
    item = _item(S.skill_metodos(ler(_artigo(tmp_path, metodos=texto))), "processo conduzido pelos autores")
    assert item["status"] == status


@pytest.mark.parametrize("texto, status", [
    ("A triagem foi realizada com auxílio do ChatGPT.", "falha"),
    ("Os dados foram extraídos por uma ferramenta automatizada.", "falha"),
    ("Um modelo de linguagem (LLM) classificou os resumos.", "falha"),
    ("Utilizou-se IA na seleção dos estudos.", "falha"),
    ("Foram selecionados estudos sobre inteligência artificial na APS.", "ok"),
    ("A média de idade foi calculada pelos autores.", "ok"),
])
def test_metodos_sem_etapa_atribuida_a_ia(tmp_path, texto, status):
    item = _item(S.skill_metodos(ler(_artigo(tmp_path, metodos=texto))), "Nenhuma etapa atribuída a IA")
    assert item["status"] == status


# ------------------------------------------------------------------ Resultados

def test_numeros_no_padrao_brasileiro(tmp_path):
    bom = _artigo(tmp_path, "bom.md",
                  resultados="Foram 1.054 participantes; OR 1,5 (IC 95% 1,2-1,9), p<0,05 e p = 0,001. "
                             "O estudo de 2019 incluiu 12.345 pessoas [3].",
                  tabela_resultados="| Nº | Achado |\n|---|---|\n| 1 | RR 0,85 (IC 95% 0,7–0,9); doi:10.1590/x.12 |")
    assert _item(S.skill_resultados(ler(bom)), "Números no padrão brasileiro")["status"] == "ok"
    ruim = _artigo(tmp_path, "ruim.md", resultados="OR 1.5 (95% CI 1.2-1.9), p<0.05.",
                   tabela_resultados="| Nº | Achado |\n|---|---|\n| 1 | RR 0.85 (CI 95% 0,7–0,9); p = 0.001 |")
    item = _item(S.skill_resultados(ler(ruim)), "Números no padrão brasileiro")
    assert item["status"] == "falha"
    for trecho in ("1.5", "95% CI", "p<0.05", "0.85", "CI 95%", "p = 0.001"):
        assert trecho in item["evidencia"]


@pytest.mark.parametrize("metodos, legendas, refs, status", [
    ("A seleção seguiu o PRISMA 2020 [2].", "", ["Silva AB. Um. Rev X. 2020;1:1.", PRISMA], "ok"),
    ("A seleção foi feita pelos autores.", "Figura 1 – Fluxograma\n\nFonte: adaptado do PRISMA 2020 [1].",
     [PRISMA], "ok"),
    ("A seleção foi feita pelos autores.", "", [PRISMA], "parcial"),
    ("A seleção seguiu o PRISMA [1].", "", ["Silva AB. Um. Rev X. 2020;1:1."], "falha"),
])
def test_modelo_prisma_2020_citado(tmp_path, metodos, legendas, refs, status):
    caminho = _artigo(tmp_path, metodos=metodos, legendas=legendas, refs=refs)
    assert _item(S.skill_resultados(ler(caminho)), "PRISMA 2020 citado")["status"] == status


# ------------------------------------------------------------------ Referências

DETALHES = ["Periódicos abreviados no padrão NLM", "Sem pontuação duplicada", "Páginas ou e-locator",
            "URL só quando", "Autores com iniciais"]


def test_referencias_bem_formatadas_nao_geram_falso_positivo(tmp_path):
    resultado = S.skill_referencias(ler(_artigo(tmp_path, refs=REFS_BOAS)))
    for criterio in DETALHES:
        assert _item(resultado, criterio)["status"] == "ok", _item(resultado, criterio)


def test_referencias_com_problemas_de_formato(tmp_path):
    resultado = S.skill_referencias(ler(_artigo(tmp_path, refs=REFS_BOAS[:2] + REFS_RUINS)))
    evidencias = {c: _item(resultado, c) for c in DETALHES}
    assert all(i["status"] != "ok" for i in evidencias.values()), evidencias
    assert "4" in evidencias["Periódicos abreviados no padrão NLM"]["evidencia"]
    assert "[4]" in evidencias["Sem pontuação duplicada"]["evidencia"]
    assert "[6]" in evidencias["Sem pontuação duplicada"]["evidencia"]
    assert "3" in evidencias["Páginas ou e-locator"]["evidencia"]
    assert "[5] URL com DOI/PMID" in evidencias["URL só quando"]["evidencia"]
    assert "[6] URL sem data de acesso" in evidencias["URL só quando"]["evidencia"]
    assert "[3] Altamirano" in evidencias["Autores com iniciais"]["evidencia"]
    # pesos baixos: os detalhes de formato não dominam a rubrica da skill
    assert sum(i["peso"] for i in evidencias.values()) <= 2.5


@pytest.mark.parametrize("ref", [
    "Silva AB, Souza CD, et al., 2020. Título qualquer. Rev X. 2020;1:1-2.",
    "SILVA, A. B., SOUZA, C. D. Título. Revista X, v. 1, p. 2, 2020.",
    "Souza AB. Reticências... no título. Rev X. 2020;1:1-2.",
])
def test_pontuacao_duplicada_ignora_et_al_iniciais_e_reticencias(ref):
    assert S._pontuacao_duplicada(ref) == []


# ------------------------------------------------------------------ comparação sistema × auditoria

def test_comparacao_com_checklist_do_sistema_inclui_itens_novos():
    externos = [{"itens": [
        {"criterio": "Métodos descrevem o processo conduzido pelos autores", "status": "ok"},
        {"criterio": S.CRITERIO_IA_METODOS, "status": "falha"},
        {"criterio": S.CRITERIO_SIGLAS, "status": "ok"},
        {"criterio": S.CRITERIO_NUMEROS, "status": "ok"},
    ]}]
    checklist = [
        {"item": "Métodos descrevem o processo conduzido pelos autores, sem atribuir etapas a ferramentas", "ok": True},
        {"item": "Siglas definidas por extenso na primeira ocorrência", "ok": True},
        {"item": "Números no padrão brasileiro (vírgula decimal)", "ok": False},
        {"item": "Item que a auditoria não confere", "ok": True},
    ]
    comp = runner._comparar_com_sistema({"checklist": checklist}, externos)
    assert [l["concorda"] for l in comp["itens"]] == [False, True, False]
    assert comp["concordancia_checklist"] == pytest.approx(33.3)


def test_artigo_bom_pontua_bem(tmp_path):
    intro = ("A hipertensão arterial sistêmica (HAS) é comum na Atenção Primária à Saúde (APS) do Sistema Único de "
             "Saúde (SUS) [1]. O objetivo desta revisão é analisar o cuidado da HAS na APS.")
    metodos = ("Dois autores buscaram, em 26/09/2026, no PubMed e no Europe PMC, com Descritores em Ciências da Saúde "
               "(DeCS) e Medical Subject Headings (MeSH) combinados por AND e OR (Quadro 1), estudos de 2016 a 2026. "
               "Os critérios de inclusão foram estudos na APS; os critérios de exclusão, editoriais. Após remover "
               "duplicatas, os autores leram título e resumo, seguindo o Preferred Reporting Items for Systematic "
               "Reviews and Meta-Analyses (PRISMA) 2020 [2].")
    quadro = ("| Base | Estratégia | Data | Registros (n) |\n|---|---|---|---|\n"
              "| PubMed | (Hypertension[MeSH Terms] AND \"Primary Health Care\"[mh]) | 26/09/2026 | 40 |")
    caminho = _artigo(tmp_path, intro=intro, metodos=metodos, tabela_metodos=quadro,
                      resultados="Foram 1.054 registros; razão de chances 1,5 (intervalo de confiança (IC) 95% 1,2-1,9) [1].",
                      discussao="Os achados divergem no SUS [1].", refs=[REFS_BOAS[1], PRISMA])
    artigo = ler(caminho)
    verificador = _VerificadorMesh()
    assert _item(S.skill_estrutura_introducao(artigo), "Siglas")["status"] == "ok"
    metodos = S.skill_metodos(artigo, verificador=verificador)
    for criterio in ("Descritores MeSH", "processo conduzido", "Nenhuma etapa atribuída"):
        assert _item(metodos, criterio)["status"] == "ok"
    resultados = S.skill_resultados(artigo)
    assert _item(resultados, "PRISMA 2020")["status"] == "ok"
    assert _item(resultados, "Números no padrão")["status"] == "ok"
    referencias = S.skill_referencias(artigo)
    assert all(_item(referencias, c)["status"] == "ok" for c in DETALHES)


def test_json_do_sistema_com_itens_novos_no_relatorio(tmp_path):
    caminho = _artigo(tmp_path, "artigo.md", intro="O SUS importa.")
    (tmp_path / "artigo.json").write_text(json.dumps({"checklist": [
        {"item": "Siglas definidas na primeira ocorrência", "ok": True}]}), encoding="utf-8")
    comp = auditar(caminho, offline=True)["comparacao_sistema"]
    assert comp["itens"] == [{"item": "Siglas definidas na primeira ocorrência", "sistema_disse_ok": True,
                              "auditoria_ok": False, "concorda": False}]
