import json
import re

import pytest
from docx import Document

from artigos_v2.pipeline import ETAPAS, Pedido, Pipeline, _problema_humanizacao, mesclar_humanizacao


def _pipeline(ambiente):
    config, ledger, cliente, fake, pubmed, epmc = ambiente
    return Pipeline(config, ledger, cliente, pubmed=pubmed, europepmc=epmc, log=lambda *_: None)


def test_pipeline_completo_gera_docx_e_registra_custos(ambiente):
    config, ledger, *_ = ambiente
    pipeline = _pipeline(ambiente)
    docx = pipeline.gerar(Pedido(tema="Hipertensão na APS", autores=["Aluno A"], orientador="Prof. B",
                                 periodo=(2016, 2026)))

    assert docx.exists() and docx.suffix == ".docx"
    artigo = ledger.artigos()[0]
    assert artigo["status"] == "concluido" and artigo["docx_path"] == str(docx)
    assert artigo["custo_api_usd"] > 0 and artigo["chamadas"] > 10

    estado = json.loads((docx.parent / "estado.json").read_text(encoding="utf-8"))
    assert estado["concluidas"] == ETAPAS
    p = estado["prisma"]
    assert p["por_base"] == {"PubMed": 40, "Europe PMC": 15}
    assert p["identificados"] - p["duplicatas"] == p["triados"] == 50
    assert p["triados"] - p["excluidos_triagem"] == p["avaliados_elegibilidade"]
    assert p["avaliados_elegibilidade"] - p["excluidos_elegibilidade"] == p["incluidos"]
    assert sum(p["motivos_triagem"].values()) == p["excluidos_triagem"]

    art = estado["artigo"]
    assert len(art["referencias"]) >= 25
    assert art["citacoes_bloqueadas"] == []
    assert len(art["sintese"]) == p["incluidos"]
    # todas as citações do texto viraram números válidos
    for texto in art["secoes"].values():
        assert not re.search(r"\[R\d+", texto)
        for grupo in re.findall(r"\{\{cite:([^}]*)\}\}", texto):
            for n in re.findall(r"\d+", grupo):
                assert 1 <= int(n) <= len(art["referencias"])

    doc = Document(str(docx))
    titulos = [par.text for par in doc.paragraphs if par.style.name.startswith("Heading")]
    assert titulos[:2] == ["RESUMO", "ABSTRACT"]
    assert "1 INTRODUÇÃO" in titulos and "5 CONCLUSÃO" in titulos and titulos[-1] == "REFERÊNCIAS"
    assert len(doc.tables) == 3  # Quadro 1 (estratégias), Figura 1 (PRISMA), Quadro 2 (síntese)
    assert any(run.font.superscript for par in doc.paragraphs for run in par.runs)
    textos = "\n".join(par.text for par in doc.paragraphs)
    assert "Orientador(a): Prof. B" in textos and "Quadro 2 – Síntese" in textos
    assert (docx.parent / "triagem.csv").exists() and (docx.parent / "artigo.md").exists()


def test_falha_no_meio_marca_erro_e_retomar_nao_refaz_etapas(ambiente):
    config, ledger, cliente, fake, *_ = ambiente
    pipeline = _pipeline(ambiente)
    artigo_id = pipeline.iniciar(Pedido(tema="Hipertensão na APS"))
    fake.explodir_em = "Escreva a DISCUSSÃO"
    with pytest.raises(Exception):
        pipeline.executar(artigo_id)
    artigo = ledger.artigo(artigo_id)
    assert artigo["status"] == "erro" and artigo["etapa_atual"] == "discussao"

    antes = len(fake.chamadas)
    fake.explodir_em = None
    docx = pipeline.executar(artigo_id)
    assert docx.exists() and ledger.artigo(artigo_id)["status"] == "concluido"
    refeitas = [c["prompt"] for c in fake.chamadas[antes:]]
    assert not any("Alinhar escopo" in p or "Triagem por TÍTULO" in p for p in refeitas)


def test_busca_refina_consulta_quando_passa_do_maximo(ambiente):
    config, ledger, cliente, fake, pubmed, epmc = ambiente
    pubmed.total = 900
    pipeline = _pipeline(ambiente)
    artigo_id = pipeline.iniciar(Pedido(tema="x", max_registros=300))
    with pytest.raises(ValueError, match="mesmo após 3 refinamentos"):
        pipeline.executar(artigo_id)
    etapas = [c["etapa"] for c in ledger.chamadas(artigo_id)]
    assert etapas.count("busca_refinamento") == 3


def test_filtros_de_data_e_idioma_entram_na_estrategia(ambiente):
    pipeline = _pipeline(ambiente)
    pedido = Pedido(tema="x", periodo=(2016, 2026), idiomas=["en", "pt"])
    assert pipeline._filtros(pedido, "pubmed", "q") == (
        '(q) AND ("2016/01/01"[dp] : "2026/12/31"[dp]) AND (english[la] OR portuguese[la]) '
        'NOT (editorial[pt] OR letter[pt] OR comment[pt])')
    assert pipeline._filtros(pedido, "europepmc", "q") == (
        '(q) AND (PUB_YEAR:[2016 TO 2026]) AND (LANG:"eng" OR LANG:"por") NOT (SRC:"PPR")')


@pytest.mark.parametrize("novo, problema", [
    ("Em 12 estudos com 340 pacientes o efeito apareceu [R1, R2]. Texto reescrito com calma.", None),
    ("Em 12 estudos com 350 pacientes o efeito apareceu [R1, R2]. Texto reescrito com calma.", "números"),
    ("Em 12 estudos com 340 pacientes o efeito apareceu [R1]. Texto reescrito com calma aqui.", "citações"),
    ("Curto.", "encurtado"),
])
def test_humanizacao_nao_pode_mudar_dados(novo, problema):
    original = "Em 12 estudos com 340 pacientes, o efeito foi observado [R1, R2]. Além disso, texto."
    resultado = _problema_humanizacao(original, novo)
    assert (resultado is None) if problema is None else (problema in resultado)


def test_humanizacao_rejeitada_mantem_texto_original(ambiente):
    config, ledger, cliente, fake, *_ = ambiente
    pipeline = _pipeline(ambiente)
    original_responder = fake.responder

    def responder(prompt):
        if "da seção DISCUSSAO" in prompt:
            return prompt.split("Texto:\n", 1)[1].replace("[R", "[R9") + " 999"
        return original_responder(prompt)

    fake.responder = responder
    docx = pipeline.gerar(Pedido(tema="x"))
    estado = json.loads((docx.parent / "estado.json").read_text(encoding="utf-8"))
    assert estado["humanizacao"]["discussao"] == "rejeitada"
    assert estado["humanizacao"]["introducao"] == "aplicada"
    assert estado["secoes"]["discussao"] == estado["secoes_originais"]["discussao"]


# ------------------------------------------------ força das afirmações (Agente de Estilo)

@pytest.mark.parametrize("original, novo, esperado", [
    ("Os dados sugerem benefício da visita domiciliar [R1].", "Os dados indicam benefício da visita domiciliar [R1].",
     "modalizador removido"),
    ("A obesidade associou-se à conversão cirúrgica [R2].", "A obesidade eleva a chance de conversão cirúrgica [R2].",
     "modalizador removido"),
    ("Uma série de casos descreveu o desfecho [R3].", "Uma coorte descreveu o desfecho em detalhe [R3].",
     "desenho de estudo alterado"),
    ("O tabagismo foi frequente entre os pacientes [R4].", "O tabagismo aumenta o risco entre os pacientes [R4].",
     "verbo de certeza"),
    ("A intervenção reduziu a pressão arterial [R5].", "A intervenção pode ter reduzido a pressão arterial [R5].",
     "mais fraca"),
    # reescritas legítimas: mesma força, "resultados" é substantivo, corrobora → confirma é troca de vocabulário
    ("Os resultados sugerem benefício, o que corrobora o achado anterior [R6].",
     "Esses resultados sugerem benefício e confirmam o achado anterior [R6].", None),
    ("A visita domiciliar associou-se a melhor adesão [R7].", "Na visita domiciliar, a adesão associou-se a ganhos [R7].",
     None),
])
def test_trava_de_forca_das_afirmacoes(original, novo, esperado):
    problema = _problema_humanizacao(original, novo)
    assert (problema is None) if esperado is None else (esperado in problema)


def test_humanizacao_por_paragrafo_preserva_so_o_que_exagerou():
    original = "Os dados sugerem benefício [R1].\n\nOutro parágrafo com texto robótico e repetitivo [R2]."
    novo = "Os dados demonstram benefício [R1].\n\nOutro parágrafo, agora com ritmo mais natural [R2]."
    final, aceitos, total, problemas = mesclar_humanizacao(original, novo)
    assert (aceitos, total) == (1, 2)
    assert final.split("\n\n") == ["Os dados sugerem benefício [R1].", "Outro parágrafo, agora com ritmo mais natural [R2]."]


# ------------------------------------------------ participação dos autores

def test_contribuicao_dos_autores_entra_nos_metodos_e_na_discussao(ambiente):
    config, ledger, cliente, fake, *_ = ambiente
    prompts_vistos = []
    responder = fake.responder
    fake.responder = lambda p: (prompts_vistos.append(p), responder(p))[1]
    contribuicao = {"triagem": "dois autores, de forma independente", "leitura_integra": "lidos na íntegra pelos autores",
                    "busca_manual": ""}
    _pipeline(ambiente).gerar(Pedido(tema="x", contribuicao_autores=contribuicao))
    metodos = next(p for p in prompts_vistos if "seção MÉTODOS" in p)
    assert "dois autores, de forma independente" in metodos and "leitura na íntegra" in metodos
    assert "busca_manual" not in metodos  # campo vazio não vira fato
    discussao = next(p for p in prompts_vistos if "Escreva a DISCUSSÃO" in p)
    assert "lidos na íntegra pelos autores" in discussao and "UM único parágrafo" in discussao


def test_autores_revisam_a_selecao_e_as_decisoes_deles_valem(ambiente):
    import csv

    config, ledger, *_ = ambiente
    pipeline = _pipeline(ambiente)
    docx = pipeline.gerar(Pedido(tema="x"))
    artigo_id = docx.parent.name
    planilha = docx.parent / "triagem.csv"
    linhas = list(csv.DictReader(planilha.open(encoding="utf-8"), delimiter=";"))
    incluidas = [l for l in linhas if l["elegibilidade"] == "incluir"]
    # autores excluem 3 estudos na leitura integral e resgatam 1 excluído na triagem
    for l in incluidas[:3]:
        l["elegibilidade"], l["motivo_elegibilidade"] = "excluir", "Desfecho fora do escopo (leitura na íntegra)"
    resgatada = next(l for l in linhas if l["triagem"] == "excluir")
    resgatada.update(triagem="incluir", motivo_triagem="", elegibilidade="incluir")
    revisada = docx.parent / "triagem-revisada.csv"
    with revisada.open("w", encoding="utf-8-sig", newline="") as f:  # como o Excel salva
        w = csv.DictWriter(f, fieldnames=linhas[0].keys(), delimiter=";")
        w.writeheader()
        w.writerows(linhas)

    antes = json.loads((docx.parent / "estado.json").read_text(encoding="utf-8"))["prisma"]
    resultado = pipeline.revisar_selecao(artigo_id, revisada)
    p = resultado["prisma"]
    assert len(resultado["alteradas"]) == 4
    assert p["incluidos"] == antes["incluidos"] - 3 + 1
    assert p["motivos_elegibilidade"]["Desfecho fora do escopo (leitura na íntegra)"] == 3
    assert ledger.artigo(artigo_id)["status"] == "em_andamento"

    docx2 = pipeline.executar(artigo_id)
    estado = json.loads((docx2.parent / "estado.json").read_text(encoding="utf-8"))
    assert len(estado["artigo"]["sintese"]) == p["incluidos"]
    assert estado["artigo"]["prisma"] == p
    assert "revisao_da_selecao" in pipeline._fatos_metodos(estado)["conduzido_pelos_autores"]


def test_revisao_com_valor_invalido_explica_o_erro(ambiente):
    import csv

    pipeline = _pipeline(ambiente)
    docx = pipeline.gerar(Pedido(tema="x"))
    linhas = list(csv.DictReader((docx.parent / "triagem.csv").open(encoding="utf-8"), delimiter=";"))
    linhas[0]["triagem"] = "talvez"
    arquivo = docx.parent / "ruim.csv"
    with arquivo.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=linhas[0].keys(), delimiter=";")
        w.writeheader()
        w.writerows(linhas)
    with pytest.raises(ValueError, match="use incluir ou excluir"):
        pipeline.revisar_selecao(docx.parent.name, arquivo)
