"""Ajustes de diagramação do .docx (revisão de design) e escape do Markdown."""

import zipfile

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

from artigos_v2.docx_export import TITULO_DECLARACAO_IA, exportar_docx, exportar_markdown

ESTRATEGIA = '("risk factor*"[tiab] OR predictor*[tiab]) AND hypertension[MeSH]'


def _artigo(**extra) -> dict:
    artigo = {
        "id": "teste", "titulo": "Hipertensão na atenção primária", "title": "Hypertension in primary care",
        "autores": ["Aluno A"], "orientador": "Prof. B", "instituicao": "Universidade X",
        "revista": "", "formato_citacao": "sobrescrito", "formato_referencias": "vancouver",
        "resumo": "Resumo do artigo.", "palavras_chave": ["Hipertensão"],
        "abstract": "Abstract text.", "keywords": ["Hypertension"],
        "secoes": {
            "introducao": "A hipertensão é frequente{{cite:1}}.",
            "metodos": "Revisão integrativa conforme o Quadro 1.",
            "resultados": "Foram incluídos dois estudos (Figura 1; Quadro 2).",
            "discussao": "Os achados convergem{{cite:1,2}}.",
            "conclusao": "Conclui-se que há lacunas.",
        },
        "estrategias": [{"base": "PubMed", "estrategia": ESTRATEGIA, "registros": 1054},
                        {"base": "Europe PMC", "estrategia": "hypertension AND primary_care", "registros": 34925}],
        "data_busca": "26/09/2026",
        "prisma": {"por_base": {"PubMed": 1054, "Europe PMC": 34925}, "identificados": 35979, "duplicatas": 979,
                   "triados": 35000, "excluidos_triagem": 34990, "motivos_triagem": {"Fora do tema": 34990},
                   "avaliados_elegibilidade": 10, "excluidos_elegibilidade": 8,
                   "motivos_elegibilidade": {"Sem desfecho": 8}, "incluidos": 2},
        "sintese": [{"autor_ano": "Silva, 2020", "citacao": "{{cite:1}}", "desenho": "Coorte",
                     "populacao": "n = 1.200 | adultos", "achados": "Fator de risco_*"},
                    {"autor_ano": "Souza, 2021", "citacao": "{{cite:2}}", "desenho": "Transversal",
                     "populacao": "300 idosos", "achados": "Associação positiva"}],
        "referencias": ["Silva A. Titulo. Rev. 2020;1:1-2.", "Souza B. Titulo. Rev. 2021;2:3-4.",
                        "Page MJ, et al. The PRISMA 2020 statement. BMJ. 2021;372:n71."],
    }
    artigo.update(extra)
    return artigo


@pytest.fixture
def docx_padrao(tmp_path):
    return exportar_docx(_artigo(), tmp_path / "artigo.docx")


def _titulos(doc):
    return [p for p in doc.paragraphs if p.style.name.startswith("Heading")]


def test_todas_as_partes_xml_sao_validas(docx_padrao):
    Document(str(docx_padrao))  # reabre com o python-docx
    with zipfile.ZipFile(docx_padrao) as pacote:
        for nome in pacote.namelist():
            if nome.endswith((".xml", ".rels")):
                etree.fromstring(pacote.read(nome))


def test_titulos_sem_numero_centralizados(docx_padrao):
    doc = Document(str(docx_padrao))
    titulos = {p.text: p for p in _titulos(doc)}
    for rotulo in ("RESUMO", "ABSTRACT", "REFERÊNCIAS"):
        assert titulos[rotulo].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert titulos["1 INTRODUÇÃO"].alignment is None  # numerados continuam como estavam
    assert TITULO_DECLARACAO_IA not in titulos


def test_quadro2_em_secao_paisagem_e_volta_a_retrato(docx_padrao):
    doc = Document(str(docx_padrao))
    assert len(doc.tables) == 3
    orientacoes = [s.orientation for s in doc.sections]
    assert orientacoes == [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]
    paisagem = doc.sections[1]
    assert paisagem.page_width > paisagem.page_height

    # o corpo é dividido pelos sectPr: o 2º trecho (seção paisagem) tem título, quadro e fonte
    trechos, atual = [], []
    for el in doc.element.body.iterchildren():
        atual.append(el)
        if el.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
            trechos.append(atual)
            atual = []
    trechos.append(atual)
    assert len(trechos) == 3
    texto = lambda el: "".join(t.text or "" for t in el.iter(qn("w:t")))
    tabelas = [el for el in trechos[1] if el.tag == qn("w:tbl")]
    assert len(tabelas) == 1 and "Principais achados" in texto(tabelas[0])
    textos = [texto(el) for el in trechos[1] if el.tag == qn("w:p")]
    assert textos[0].startswith("Quadro 2 –") and any(t.startswith("Fonte:") for t in textos)
    assert "2 MÉTODOS" in [texto(el) for el in trechos[0]] and "4 DISCUSSÃO" in [texto(el) for el in trechos[2]]

    # larguras do Quadro 2 preenchem a largura útil da página deitada
    util = paisagem.page_width - paisagem.left_margin - paisagem.right_margin
    grade = tabelas[0].find(qn("w:tblGrid"))
    soma = sum(int(c.get(qn("w:w"))) for c in grade.iterchildren()) * 635  # twips → EMU
    assert abs(soma - util) < 36000  # < 1 mm


def test_tabelas_com_layout_fixo_e_cabecalho_repetido(docx_padrao):
    doc = Document(str(docx_padrao))
    for tabela in doc.tables:
        layout = tabela._tbl.tblPr.find(qn("w:tblLayout"))
        assert layout is not None and layout.get(qn("w:type")) == "fixed"
        assert all(int(c.get(qn("w:w"))) > 0 for c in tabela._tbl.tblGrid.iterchildren())
    quadro1, _, quadro2 = doc.tables
    for quadro in (quadro1, quadro2):
        assert quadro.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    # colunas estreitas (Nº, Data, Registros) mantêm a largura mínima
    assert quadro2.rows[1].cells[0].width.cm >= 0.99
    assert quadro1.rows[1].cells[2].width.cm >= 2.29 and quadro1.rows[1].cells[3].width.cm >= 2.09


def test_legendas_em_10pt_e_titulo_preso_ao_objeto(docx_padrao):
    doc = Document(str(docx_padrao))
    legendas = [p for p in doc.paragraphs if p.text.startswith(("Quadro ", "Figura ", "Fonte:"))]
    assert len(legendas) == 6
    for p in legendas:
        assert all(r.font.size.pt == 10 for r in p.runs)
        if not p.text.startswith("Fonte:"):
            assert p.paragraph_format.keep_with_next and all(r.font.bold for r in p.runs)


def test_hifenizacao_e_idioma_pt_br(docx_padrao):
    doc = Document(str(docx_padrao))
    hifen = doc.settings.element.find(qn("w:autoHyphenation"))
    assert hifen is not None and hifen.get(qn("w:val")) == "true"
    padrao = doc.styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr/w:lang")
    assert padrao and padrao[0].get(qn("w:val")) == "pt-BR"
    normal = doc.styles["Normal"].element.rPr.find(qn("w:lang"))
    assert normal.get(qn("w:val")) == "pt-BR"
    # a fonte do título não fica presa ao tema (Calibri/Cambria)
    fontes = doc.styles["Heading 1"].element.rPr.rFonts
    assert fontes.get(qn("w:asciiTheme")) is None and fontes.get(qn("w:ascii")) == "Times New Roman"


def test_prisma_com_n_inquebravel_e_milhar(docx_padrao):
    doc = Document(str(docx_padrao))
    prisma = doc.tables[1]
    textos = [p for linha in prisma.rows for c in linha.cells for p in c.paragraphs]
    tudo = "\n".join(p.text for p in textos)
    assert "n = " in tudo and "(n = " not in tudo
    assert "PubMed (n = 1.054)" in tudo and "34.925" in tudo and "35.979" in tudo
    assert all(r.font.size.pt == 9 for p in textos for r in p.runs if r.text.strip() not in ("", "→", "↓"))
    registros = [linha.cells[3].text for linha in doc.tables[0].rows[1:]]
    assert registros == ["1.054", "34.925"]


def test_fonte_da_figura_sem_citacao_mantem_texto(docx_padrao):
    doc = Document(str(docx_padrao))
    assert any(p.text.startswith("Fonte: elaborado pelos autores com base no modelo PRISMA 2020")
               for p in doc.paragraphs)


def test_prisma_citacao_vira_sobrescrito_na_fonte(tmp_path):
    doc = Document(str(exportar_docx(_artigo(prisma_citacao="{{cite:3}}"), tmp_path / "a.docx")))
    fonte = next(p for p in doc.paragraphs if p.text.startswith("Fonte: adaptado de Page et al."))
    assert fonte.text == "Fonte: adaptado de Page et al.3 (modelo PRISMA 2020)."
    assert [r.text for r in fonte.runs if r.font.superscript] == ["3"]
    assert all(r.font.size.pt == 10 for r in fonte.runs)


def test_declaracao_ia_so_quando_informada(tmp_path):
    declaracao = "Os autores usaram ferramenta de IA para revisão de linguagem e revisaram todo o texto."
    doc = Document(str(exportar_docx(_artigo(declaracao_ia=declaracao), tmp_path / "a.docx")))
    titulos = [p.text for p in _titulos(doc)]
    assert titulos.index("5 CONCLUSÃO") < titulos.index(TITULO_DECLARACAO_IA) == len(titulos) - 2
    assert titulos[-1] == "REFERÊNCIAS"
    titulo = next(p for p in _titulos(doc) if p.text == TITULO_DECLARACAO_IA)
    assert titulo.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert declaracao in [p.text for p in doc.paragraphs]

    sem = Document(str(exportar_docx(_artigo(declaracao_ia=""), tmp_path / "b.docx")))
    texto = "\n".join(p.text for p in sem.paragraphs)
    assert TITULO_DECLARACAO_IA not in texto and "inteligência artificial" not in texto.lower()


def test_estrategia_com_asterisco_preservada_no_docx(docx_padrao):
    doc = Document(str(docx_padrao))
    assert doc.tables[0].rows[1].cells[1].text == ESTRATEGIA


def test_markdown_escapa_curingas_e_celulas(tmp_path):
    md = exportar_markdown(_artigo(declaracao_ia="Declaração curta."), tmp_path / "a.md").read_text(encoding="utf-8")
    assert r'("risk factor\*"[tiab] OR predictor\*[tiab])' in md
    assert r"primary\_care" in md
    assert r"n = 1.200 \| adultos" in md and r"Fator de risco\_\*" in md
    assert "predictor*[" not in md  # nenhum asterisco solto que vire itálico
    assert "## Declaração de uso de inteligência artificial" in md
    assert md.index("## Conclusão") < md.index("## Declaração") < md.index("## Referências")
    sem = exportar_markdown(_artigo(), tmp_path / "b.md").read_text(encoding="utf-8")
    assert "Declaração de uso" not in sem
