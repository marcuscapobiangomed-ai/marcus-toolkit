"""Exporta o artigo final (artigo.json) para .docx e .md.

Chamado automaticamente na última etapa do pipeline; também disponível via
`python -m artigos_v2 exportar <id>` para reexportar sem gastar IA.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .referencias import renderizar_citacoes, texto_plano

TITULOS = [("introducao", "Introdução"), ("metodos", "Métodos"), ("resultados", "Resultados"),
           ("discussao", "Discussão"), ("conclusao", "Conclusão")]
FONTE = "Times New Roman"


def _configurar(documento: Document) -> None:
    secao = documento.sections[0]
    secao.orientation = WD_ORIENT.PORTRAIT
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    secao.top_margin = secao.left_margin = Cm(3)
    secao.bottom_margin = secao.right_margin = Cm(2)

    normal = documento.styles["Normal"]
    _fonte(normal)
    normal.font.size = Pt(12)
    formato = normal.paragraph_format
    formato.line_spacing = 1.5
    formato.space_after = Pt(6)

    titulo = documento.styles["Heading 1"]
    _fonte(titulo)
    titulo.font.color.rgb = RGBColor(0, 0, 0)
    documento.styles["Heading 1"].font.size = Pt(12)
    documento.styles["Heading 1"].font.bold = True
    documento.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    documento.styles["Heading 1"].paragraph_format.space_after = Pt(6)

    # Numeração de página no rodapé
    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _campo(rodape.add_run(), "PAGE")


def _fonte(estilo) -> None:
    estilo.font.name = FONTE
    estilo.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONTE)


def _campo(run, instrucao: str) -> None:
    for tipo, texto in (("begin", None), (None, instrucao), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def _paragrafo(documento, texto: str, formato_citacao: str, recuo: bool = True, tamanho: int | None = None,
               alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = documento.add_paragraph()
    p.alignment = alinhamento
    if recuo:
        p.paragraph_format.first_line_indent = Cm(1.25)
    for trecho, sobrescrito in renderizar_citacoes(texto, formato_citacao):
        run = p.add_run(trecho)
        run.font.superscript = sobrescrito
        if tamanho:
            run.font.size = Pt(tamanho)
    return p


def _texto_em_celula(celula, texto: str, formato_citacao: str = "sobrescrito", negrito=False, tamanho=10,
                     alinhamento=WD_ALIGN_PARAGRAPH.LEFT):
    celula.text = ""
    linhas = texto.split("\n")
    for i, linha in enumerate(linhas):
        p = celula.paragraphs[0] if i == 0 else celula.add_paragraph()
        p.alignment = alinhamento
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Cm(0)
        for trecho, sobrescrito in renderizar_citacoes(linha, formato_citacao):
            run = p.add_run(trecho)
            run.font.size = Pt(tamanho)
            run.font.bold = negrito
            run.font.superscript = sobrescrito


def _bordas(celula, visivel: bool = True) -> None:
    tc_pr = celula._tc.get_or_add_tcPr()
    bordas = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single" if visivel else "nil")
        if visivel:
            el.set(qn("w:sz"), "8")
            el.set(qn("w:color"), "000000")
        bordas.append(el)
    tc_pr.append(bordas)


def _sombrear(celula, cor: str) -> None:
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(sombra)


def _legenda(documento, texto: str, antes: bool = True):
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if antes else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = antes
    run = p.add_run(texto)
    run.font.size = Pt(10)
    run.font.bold = antes
    return p


def _tabela(documento, cabecalho: list[str], linhas: list[list[str]], larguras_cm: list[float],
            formato_citacao: str):
    tabela = documento.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, titulo in enumerate(cabecalho):
        _texto_em_celula(tabela.rows[0].cells[i], titulo, negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
        _sombrear(tabela.rows[0].cells[i], "D9D9D9")
    for linha in linhas:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha):
            _texto_em_celula(celulas[i], valor, formato_citacao)
    for linha in tabela.rows:
        for i, largura in enumerate(larguras_cm):
            linha.cells[i].width = Cm(largura)
    # repete o cabeçalho em quebras de página
    tr_pr = tabela.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    tr_pr.append(repetir)
    return tabela


def _fluxograma_prisma(documento, prisma: dict) -> None:
    """Fluxograma PRISMA como tabela com caixas — editável no Word e sem dependência gráfica."""
    por_base = "; ".join(f"{base} (n = {n})" for base, n in prisma["por_base"].items())
    motivos = lambda m: "\n".join(f"• {motivo} (n = {n})" for motivo, n in sorted(m.items(), key=lambda x: -x[1]))
    linhas = [
        ("Identificação", f"Registros identificados nas bases de dados (n = {prisma['identificados']})\n{por_base}",
         f"Duplicatas removidas (n = {prisma['duplicatas']})"),
        ("Triagem", f"Registros triados por título e resumo (n = {prisma['triados']})",
         f"Registros excluídos (n = {prisma['excluidos_triagem']})\n{motivos(prisma['motivos_triagem'])}".strip()),
        ("Elegibilidade", f"Artigos avaliados para elegibilidade (n = {prisma['avaliados_elegibilidade']})",
         f"Artigos excluídos (n = {prisma['excluidos_elegibilidade']})\n"
         f"{motivos(prisma['motivos_elegibilidade'])}".strip()),
        ("Inclusão", f"Estudos incluídos na revisão (n = {prisma['incluidos']})", ""),
    ]
    tabela = documento.add_table(rows=0, cols=4)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (etapa, principal, lateral) in enumerate(linhas):
        c = tabela.add_row().cells
        _texto_em_celula(c[0], etapa, negrito=True, tamanho=9, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
        _sombrear(c[0], "D9E2F3")
        _bordas(c[0])
        _texto_em_celula(c[1], principal, tamanho=9, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
        _bordas(c[1])
        _texto_em_celula(c[2], "→" if lateral else "", tamanho=12, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
        _bordas(c[2], visivel=False)
        _texto_em_celula(c[3], lateral, tamanho=9)
        _bordas(c[3], visivel=bool(lateral))
        if i < len(linhas) - 1:
            seta = tabela.add_row().cells
            for j, celula in enumerate(seta):
                _texto_em_celula(celula, "↓" if j == 1 else "", tamanho=12, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
                _bordas(celula, visivel=False)
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, (2.6, 6.2, 0.8, 6.4)):
            celula.width = Cm(largura)


def exportar_docx(artigo: dict, destino: str | Path) -> Path:
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    fc = artigo.get("formato_citacao", "sobrescrito")
    ano = date.today().year
    documento = Document()
    _configurar(documento)

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(artigo["titulo"].upper())
    run.bold = True
    if artigo.get("title"):
        t = documento.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.add_run(artigo["title"]).italic = True

    if artigo.get("autores"):
        _paragrafo(documento, ", ".join(artigo["autores"]), fc, recuo=False, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    if artigo.get("orientador"):
        _paragrafo(documento, f"Orientador(a): {artigo['orientador']}", fc, recuo=False,
                   alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    if artigo.get("instituicao"):
        _paragrafo(documento, artigo["instituicao"], fc, recuo=False, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    for rotulo, texto, rotulo_chave, chaves in (
        ("RESUMO", artigo.get("resumo"), "Palavras-chave", artigo.get("palavras_chave")),
        ("ABSTRACT", artigo.get("abstract"), "Keywords", artigo.get("keywords")),
    ):
        if not texto:
            continue
        documento.add_heading(rotulo, level=1)
        p = _paragrafo(documento, texto, fc, recuo=False)
        p.paragraph_format.line_spacing = 1.0
        if chaves:
            p = documento.add_paragraph()
            p.add_run(f"{rotulo_chave}: ").bold = True
            p.add_run("; ".join(chaves) + ".")

    numero = 0
    for chave, rotulo in TITULOS:
        texto = artigo["secoes"].get(chave, "")
        if not texto:
            continue
        numero += 1
        documento.add_heading(f"{numero} {rotulo.upper()}", level=1)
        for bloco in [b.strip() for b in texto.split("\n\n") if b.strip()]:
            _paragrafo(documento, " ".join(bloco.splitlines()), fc)

        if chave == "metodos":
            _legenda(documento, "Quadro 1 – Estratégias de busca utilizadas em cada base de dados")
            _tabela(documento, ["Base de dados", "Estratégia de busca", "Data da busca", "Registros (n)"],
                    [[e["base"], e["estrategia"], artigo.get("data_busca", ""), str(e["registros"])]
                     for e in artigo["estrategias"]],
                    [2.8, 9.4, 2.3, 1.9], fc)
            _legenda(documento, f"Fonte: elaborado pelos autores ({ano}).", antes=False)

        if chave == "resultados":
            _legenda(documento, "Figura 1 – Fluxograma do processo de identificação, triagem e inclusão dos "
                                "estudos (modelo PRISMA)")
            _fluxograma_prisma(documento, artigo["prisma"])
            _legenda(documento, f"Fonte: elaborado pelos autores com base no modelo PRISMA 2020 ({ano}).",
                     antes=False)

            _legenda(documento, "Quadro 2 – Síntese dos estudos incluídos na revisão")
            _tabela(documento, ["Nº", "Autor/ano", "Desenho do estudo", "População/amostra", "Principais achados"],
                    [[str(i + 1), f"{s['autor_ano']}{s['citacao']}", s["desenho"], s["populacao"], s["achados"]]
                     for i, s in enumerate(artigo["sintese"])],
                    [0.9, 3.0, 3.0, 3.6, 5.9], fc)
            _legenda(documento, f"Fonte: elaborado pelos autores ({ano}).", antes=False)

    documento.add_heading("REFERÊNCIAS", level=1)
    for i, referencia in enumerate(artigo["referencias"], start=1):
        p = documento.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(8)
        prefixo = f"{i}. " if artigo.get("formato_referencias", "vancouver") == "vancouver" else ""
        p.add_run(prefixo + referencia)

    documento.core_properties.title = artigo["titulo"]
    documento.core_properties.author = ", ".join(artigo.get("autores") or [])
    documento.save(destino)
    return destino


def exportar_markdown(artigo: dict, destino: str | Path) -> Path:
    destino = Path(destino)
    linhas = [f"# {artigo['titulo']}", ""]
    if artigo.get("resumo"):
        linhas += ["## Resumo", artigo["resumo"], "", "**Palavras-chave:** " + "; ".join(artigo.get("palavras_chave", [])), ""]
    for chave, rotulo in TITULOS:
        if artigo["secoes"].get(chave):
            linhas += [f"## {rotulo}", texto_plano(artigo["secoes"][chave], "colchetes"), ""]
        if chave == "resultados":
            p = artigo["prisma"]
            linhas += ["**Fluxograma PRISMA:** " + ", ".join(f"{b}: {n}" for b, n in p["por_base"].items())
                       + f" → identificados {p['identificados']} → duplicatas {p['duplicatas']} → triados "
                         f"{p['triados']} (excluídos {p['excluidos_triagem']}) → elegibilidade "
                         f"{p['avaliados_elegibilidade']} (excluídos {p['excluidos_elegibilidade']}) → incluídos "
                         f"{p['incluidos']}", ""]
            linhas += ["| Nº | Autor/ano | Desenho | População | Achados |", "|---|---|---|---|---|"]
            linhas += [f"| {i + 1} | {s['autor_ano']} {texto_plano(s['citacao'])} | {s['desenho']} | "
                       f"{s['populacao']} | {s['achados']} |" for i, s in enumerate(artigo["sintese"])]
            linhas.append("")
    linhas += ["## Referências"] + [f"{i}. {r}" for i, r in enumerate(artigo["referencias"], 1)]
    destino.write_text("\n".join(linhas) + "\n", encoding="utf-8")
    return destino
