"""Leitura de .docx / .md / .txt em seções, tabelas, legendas e citações."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

SECOES = {
    "resumo": r"resumo",
    "abstract": r"abstract",
    "introducao": r"introdu[cç][aã]o",
    "metodos": r"(materia(?:l|is) e )?m[eé]todos?|metodologia|percurso metodol[oó]gico",
    "resultados_discussao": r"resultados e discuss[aã]o",
    "resultados": r"resultados",
    "discussao": r"discuss[aã]o",
    "conclusao": r"conclus(?:[aã]o|[oõ]es)|considera[cç][oõ]es finais",
    "referencias": r"refer[eê]ncias(?: bibliogr[aá]ficas)?|bibliografia",
}


@dataclass
class Tabela:
    linhas: list[list[str]]
    secao: str

    @property
    def texto(self) -> str:
        return "\n".join(" | ".join(c for c in linha) for linha in self.linhas)


@dataclass
class Artigo:
    caminho: str
    secoes: dict[str, list[str]] = field(default_factory=dict)  # seção -> parágrafos
    tabelas: list[Tabela] = field(default_factory=list)
    legendas: list[str] = field(default_factory=list)
    referencias: list[str] = field(default_factory=list)
    preambulo: list[str] = field(default_factory=list)
    # subtítulos dentro das seções (### 2.1 ..., estilo Título 2): não fazem parte do texto da seção
    subtitulos: list[str] = field(default_factory=list)
    # citações numéricas encontradas no corpo (texto sobrescrito ou entre colchetes), em ordem
    citacoes: list[list[int]] = field(default_factory=list)

    def texto(self, secao: str) -> str:
        return "\n\n".join(self.secoes.get(secao, []))

    @property
    def corpo(self) -> str:
        return "\n\n".join(self.texto(s) for s in ("introducao", "metodos", "resultados", "resultados_discussao",
                                                   "discussao", "conclusao"))


def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def identificar_secao(texto: str) -> str | None:
    limpo = re.sub(r"^\s*(\d+(\.\d+)*[\.\)]?|[IVX]+[\.\)])\s*", "", texto.strip()).strip(" :.").lower()
    if len(limpo.split()) > 6:
        return None
    for nome, padrao in SECOES.items():
        if re.fullmatch(padrao, limpo) or re.fullmatch(padrao, _sem_acento(limpo)):
            return nome
    return None


def expandir_numeros(trecho: str) -> list[int]:
    numeros = []
    for parte in re.split(r"[,;]\s*", trecho):
        faixa = re.fullmatch(r"\s*(\d+)\s*[-–—]\s*(\d+)\s*", parte)
        if faixa:
            a, b = int(faixa.group(1)), int(faixa.group(2))
            if a <= b <= a + 60:
                numeros.extend(range(a, b + 1))
        elif re.fullmatch(r"\s*\d+\s*", parte):
            numeros.append(int(parte))
    return numeros


def _eh_legenda(texto: str) -> bool:
    return bool(re.match(r"^\s*(figura|quadro|tabela|gr[aá]fico|fonte)\b", texto, re.I))


def ler(caminho: str | Path) -> Artigo:
    caminho = Path(caminho)
    if caminho.suffix.lower() == ".docx":
        return _ler_docx(caminho)
    return _ler_texto(caminho)


def _ler_docx(caminho: Path) -> Artigo:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    documento = Document(str(caminho))
    artigo = Artigo(str(caminho))
    atual = None

    for elemento in documento.element.body.iterchildren():
        if elemento.tag.endswith("}tbl"):
            tabela = Table(elemento, documento)
            linhas = []
            for linha in tabela.rows:
                celulas = []
                for celula in linha.cells:
                    celulas.append(" ".join(_texto_paragrafo(p, artigo, contar=False)[0] for p in celula.paragraphs).strip())
                # células mescladas se repetem; remove repetição consecutiva
                linhas.append([c for i, c in enumerate(celulas) if i == 0 or c != celulas[i - 1]])
            artigo.tabelas.append(Tabela(linhas, atual or "preambulo"))
            # citações sobrescritas em tabelas (quadro-síntese) também contam
            for linha in tabela.rows:
                for celula in linha.cells:
                    for p in celula.paragraphs:
                        _texto_paragrafo(p, artigo, contar=True)
            continue
        if not elemento.tag.endswith("}p"):
            continue
        paragrafo = Paragraph(elemento, documento)
        estilo = (paragrafo.style.name or "").lower() if paragrafo.style is not None else ""
        bruto = paragrafo.text.strip()
        if not bruto:
            continue
        secao = identificar_secao(bruto)
        cabecalho = secao and ("heading" in estilo or "título" in estilo or "titulo" in estilo
                               or all(r.bold for r in paragrafo.runs if r.text.strip()) or bruto.isupper()
                               or len(bruto.split()) <= 4)
        if cabecalho:
            atual = secao
            artigo.secoes.setdefault(atual, [])
            continue
        if atual not in (None, "referencias") and ("heading" in estilo or "título" in estilo or "titulo" in estilo):
            artigo.subtitulos.append(bruto)
            continue
        if atual == "referencias":  # antes da legenda: uma referência pode começar por "Fonte AB."
            artigo.referencias.append(bruto)
            continue
        if _eh_legenda(bruto):
            # a linha "Fonte:" pode citar referência (ex.: modelo PRISMA 2020); a citação conta
            texto, _ = _texto_paragrafo(paragrafo, artigo, contar=atual not in (None, "resumo", "abstract"))
            artigo.legendas.append(texto)
            continue
        texto, _ = _texto_paragrafo(paragrafo, artigo, contar=atual not in (None, "resumo", "abstract"))
        if atual is None:
            artigo.preambulo.append(texto)
        else:
            artigo.secoes[atual].append(texto)
    _dividir_resultados_discussao(artigo)
    return artigo


def _texto_paragrafo(paragrafo, artigo: Artigo, contar: bool) -> tuple[str, list]:
    """Texto do parágrafo com números sobrescritos marcados como [^1,2]; registra as citações."""
    partes, sobrescrito_atual = [], ""
    for run in paragrafo.runs:
        if run.font.superscript and re.fullmatch(r"[\d,;\s\-–—]+", run.text or ""):
            sobrescrito_atual += run.text
            continue
        if sobrescrito_atual:
            partes.append(f"[^{sobrescrito_atual.strip()}]")
            sobrescrito_atual = ""
        partes.append(run.text)
    if sobrescrito_atual:
        partes.append(f"[^{sobrescrito_atual.strip()}]")
    texto = "".join(partes).strip()
    if contar:
        for m in re.finditer(r"\[\^([^\]]+)\]|\[(\d+(?:\s*[,;\-–]\s*\d+)*)\]", texto):
            numeros = expandir_numeros(m.group(1) or m.group(2))
            if numeros:
                artigo.citacoes.append(numeros)
    return texto, []


def _ler_texto(caminho: Path) -> Artigo:
    artigo = Artigo(str(caminho))
    atual = None
    blocos = re.split(r"\n\s*\n", caminho.read_text(encoding="utf-8", errors="replace"))
    tabela_atual: list[list[str]] = []
    for bloco in blocos:
        for linha in bloco.splitlines():
            if linha.strip().startswith("|"):
                if not re.fullmatch(r"\s*\|[\s\-:|]+\|\s*", linha):
                    tabela_atual.append([c.strip() for c in linha.strip().strip("|").split("|")])
        if tabela_atual:
            artigo.tabelas.append(Tabela(tabela_atual, atual or "preambulo"))
            tabela_atual = []
            continue
        texto = " ".join(l.strip() for l in bloco.splitlines()).strip()
        if not texto:
            continue
        cabecalho = re.match(r"^#{1,6}\s*(.+)$", texto)
        secao = identificar_secao(cabecalho.group(1) if cabecalho else texto)
        if secao:
            atual = secao
            artigo.secoes.setdefault(atual, [])
            continue
        if cabecalho and atual not in (None, "referencias"):
            artigo.subtitulos.append(cabecalho.group(1).strip())
            continue
        if atual == "referencias":
            artigo.referencias.extend(l.strip() for l in bloco.splitlines() if l.strip())
            continue
        if _eh_legenda(texto):
            artigo.legendas.append(texto)
            if atual not in (None, "resumo", "abstract"):
                for m in re.finditer(r"\[(\d+(?:\s*[,;\-–]\s*\d+)*)\]", texto):
                    artigo.citacoes.append(expandir_numeros(m.group(1)))
            continue
        if atual in (None,):
            artigo.preambulo.append(texto)
            continue
        artigo.secoes[atual].append(texto)
        if atual not in ("resumo", "abstract"):
            for m in re.finditer(r"\[(\d+(?:\s*[,;\-–]\s*\d+)*)\]", texto):
                artigo.citacoes.append(expandir_numeros(m.group(1)))
    _dividir_resultados_discussao(artigo)
    return artigo


def _dividir_resultados_discussao(artigo: Artigo) -> None:
    """'Resultados e discussão' conta para as duas seções."""
    if "resultados_discussao" in artigo.secoes:
        texto = artigo.secoes.pop("resultados_discussao")
        artigo.secoes.setdefault("resultados", []).extend(texto)
        artigo.secoes.setdefault("discussao", []).extend(texto)
        for t in artigo.tabelas:
            if t.secao == "resultados_discussao":
                t.secao = "resultados"
